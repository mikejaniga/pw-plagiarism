"""Generate a multi-sheet Excel report (.xlsx) with coloring and comparative analysis."""

from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import openpyxl
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from plagiarism_detector.data_loader import Document
from plagiarism_detector.models.ensemble import PairwiseDocResult, SentenceMatch


class ExcelPlagiarismReporter:
    def __init__(self, model_1_name: str, model_2_name: str, alpha: float) -> None:
        self.model_1_name: str = model_1_name
        self.model_2_name: str = model_2_name
        self.alpha: float = alpha

    def generate_report(
        self,
        output_path: Path,
        documents: List[Document],
        pairwise_results: List[PairwiseDocResult],
        sentence_matches: List[SentenceMatch],
        matrices: Dict[str, np.ndarray],
    ) -> None:
        """Generate a comprehensive Excel spreadsheet."""
        wb: openpyxl.Workbook = openpyxl.Workbook()
        # Default sheet
        wb.remove(wb.active)

        self._create_summary_sheet(wb, pairwise_results)
        self._create_matrices_sheet(wb, documents, matrices)
        self._create_sentence_matches_sheet(wb, sentence_matches)
        self._create_model_comparison_sheet(wb, pairwise_results, sentence_matches)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        print(f"-> Excel report has been saved to: {output_path}")

    def _style_header_row(
        self, ws: Worksheet, row_idx: int, max_col: int, bg_color: str = "1F497D"
    ) -> None:
        """Formats the table header."""
        header_fill: PatternFill = PatternFill(
            start_color=bg_color, end_color=bg_color, fill_type="solid"
        )
        header_font: Font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border: Border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="medium", color="1F497D"),
        )
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = thin_border
        ws.row_dimensions[row_idx].height = 28

    def _auto_fit_columns(self, ws: Worksheet, max_len_cap: int = 60) -> None:
        """Adjusts column widths."""
        for col in ws.columns:
            max_len: int = 0
            col_letter: str = get_column_letter(col[0].column)
            for cell in col:
                val: str = str(cell.value or "")
                # Take the first line in case of multiline
                val_first: str = val.split("\n")[0]
                max_len = max(max_len, len(val_first))
            ws.column_dimensions[col_letter].width = min(
                max(max_len + 3, 12), max_len_cap
            )

    def _create_summary_sheet(
        self, wb: openpyxl.Workbook, results: List[PairwiseDocResult]
    ) -> None:
        """Creates Sheet 1: Pairwise comparison summary."""
        ws: Worksheet = wb.create_sheet(title="Pair Summary")
        ws.views.sheetView[0].showGridLines = True

        # Title
        ws.merge_cells("A1:I1")
        title_cell = ws["A1"]
        title_cell.value = "PLAGIARISM DETECTION REPORT - PAIRWISE DOCUMENT COMPARISON"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F497D")
        ws.row_dimensions[1].height = 30

        headers: List[str] = [
            "Document A",
            "Document B",
            f"Similarity: {self.model_1_name}",
            f"Similarity: {self.model_2_name}",
            f"Ensemble (weight {self.alpha:.2f})",
            "Coverage Score",
            "Number of Detected Sentences",
            "Mean Sentence Similarity",
            "Plagiarism Risk Assessment",
        ]

        ws.append([])
        ws.append(headers)
        self._style_header_row(ws, row_idx=3, max_col=len(headers), bg_color="1F497D")

        for r in results:
            row_data = [
                r.doc_a_name,
                r.doc_b_name,
                round(r.doc_sim_model_1, 4),
                round(r.doc_sim_model_2, 4),
                round(r.doc_sim_ensemble, 4),
                round(r.coverage_score, 4),
                r.matched_sentences_count,
                round(r.mean_match_similarity, 4)
                if r.matched_sentences_count > 0
                else 0.0,
                r.risk_level,
            ]
            ws.append(row_data)

        # Number formatting and risk coloring
        for row in ws.iter_rows(
            min_row=4, max_row=3 + len(results), min_col=1, max_col=9
        ):
            for i, cell in enumerate(row):
                cell.alignment = Alignment(
                    vertical="center", horizontal="center" if i >= 2 else "left"
                )
                if i in [2, 3, 4, 5, 7]:
                    cell.number_format = "0.00%"
                # Risk color marking
                if i == 8:
                    val: str = str(cell.value)
                    if "Very high" in val:
                        cell.fill = PatternFill("solid", fgColor="FADBD8")
                        cell.font = Font(color="900C3F", bold=True)
                    elif "High" in val:
                        cell.fill = PatternFill("solid", fgColor="FCF3CF")
                        cell.font = Font(color="9A7D0A", bold=True)
                    elif "Moderate" in val:
                        cell.fill = PatternFill("solid", fgColor="E8F8F5")
                        cell.font = Font(color="117864")

        self._auto_fit_columns(ws)

    def _create_matrices_sheet(
        self,
        wb: openpyxl.Workbook,
        documents: List[Document],
        matrices: Dict[str, np.ndarray],
    ) -> None:
        """Creates Sheet 2: Document similarity matrices for both models and the ensemble."""
        ws: Worksheet = wb.create_sheet(title="Similarity Matrices")
        ws.views.sheetView[0].showGridLines = True

        doc_names: List[str] = [d.file_path.name for d in documents]
        n: int = len(doc_names)

        sections = [
            ("Model 1 Matrix: " + self.model_1_name, matrices["model_1"]),
            ("Model 2 Matrix: " + self.model_2_name, matrices["model_2"]),
            (f"Hybrid Matrix (Ensemble alpha={self.alpha})", matrices["ensemble"]),
        ]

        current_row: int = 1
        for title, mat in sections:
            ws.cell(row=current_row, column=1, value=title).font = Font(
                size=12, bold=True, color="1F497D"
            )
            current_row += 1

            # Column headers
            headers = ["Document"] + doc_names
            ws.row_dimensions[current_row].height = 24
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=current_row, column=c_idx, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="34495E")
                cell.alignment = Alignment(horizontal="center", vertical="center")

            start_data_row: int = current_row + 1
            for r_idx in range(n):
                row_cells = [doc_names[r_idx]] + [
                    round(float(mat[r_idx, c_idx]), 4) for c_idx in range(n)
                ]
                row_num: int = start_data_row + r_idx
                for c_idx, val in enumerate(row_cells, start=1):
                    cell = ws.cell(row=row_num, column=c_idx, value=val)
                    if c_idx > 1:
                        cell.number_format = "0.0%"
                        cell.alignment = Alignment(horizontal="center")
                    else:
                        cell.font = Font(bold=True)

            end_data_row: int = start_data_row + n - 1

            # Three-color scale (White -> Yellow -> Red)
            rule: ColorScaleRule = ColorScaleRule(
                start_type="num",
                start_value=0.2,
                start_color="FFFFFF",
                mid_type="num",
                mid_value=0.6,
                mid_color="F9E79F",
                end_type="num",
                end_value=1.0,
                end_color="E74C3C",
            )
            col_start_letter: str = get_column_letter(2)
            col_end_letter: str = get_column_letter(1 + n)
            ws.conditional_formatting.add(
                f"{col_start_letter}{start_data_row}:{col_end_letter}{end_data_row}",
                rule,
            )

            current_row = end_data_row + 3

        self._auto_fit_columns(ws)

    def _create_sentence_matches_sheet(
        self, wb: openpyxl.Workbook, matches: List[SentenceMatch]
    ) -> None:
        """Creates Sheet 3: Detailed listing of suspicious sentence fragments."""
        ws: Worksheet = wb.create_sheet(title="Detected Fragments")
        ws.views.sheetView[0].showGridLines = True

        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = (
            "DETAILED ANALYSIS OF SUSPICIOUS SENTENCE PAIRS (PARAPHRASES AND COPIES)"
        )
        title_cell.font = Font(name="Calibri", size=13, bold=True, color="1F497D")

        headers: List[str] = [
            "Document A",
            "Fragment from Document A",
            "Document B",
            "Fragment from Document B",
            f"Score {self.model_1_name}",
            f"Score {self.model_2_name}",
            "Ensemble Score",
            "Divergence |M1 - M2|",
        ]

        ws.append([])
        ws.append(headers)
        self._style_header_row(ws, row_idx=3, max_col=len(headers), bg_color="2E4053")

        for m in matches:
            ws.append(
                [
                    m.doc_a_name,
                    m.sent_a_text,
                    m.doc_b_name,
                    m.sent_b_text,
                    round(m.score_model_1, 4),
                    round(m.score_model_2, 4),
                    round(m.score_ensemble, 4),
                    round(m.divergence, 4),
                ]
            )

        # Formatting
        for row in ws.iter_rows(
            min_row=4, max_row=3 + len(matches), min_col=1, max_col=8
        ):
            for i, cell in enumerate(row):
                if i in [1, 3]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                elif i >= 4:
                    cell.number_format = "0.0%"
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                else:
                    cell.alignment = Alignment(vertical="top")

        # Width settings
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 50
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 16
        ws.column_dimensions["G"].width = 16
        ws.column_dimensions["H"].width = 16

    def _create_model_comparison_sheet(
        self,
        wb: openpyxl.Workbook,
        pairwise_results: List[PairwiseDocResult],
        matches: List[SentenceMatch],
    ) -> None:
        """Creates Sheet 4: Model comparison and analysis of the benefits of fusion."""
        ws: Worksheet = wb.create_sheet(title="Model Evaluation")
        ws.views.sheetView[0].showGridLines = True

        ws.merge_cells("A1:F1")
        ws["A1"].value = "STUDY OF THE COMBINATION OF TWO LANGUAGE MODELS (ENSEMBLE MODEL)"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F497D")

        # Statistical calculations
        m1_scores: List[float] = [r.doc_sim_model_1 for r in pairwise_results]
        m2_scores: List[float] = [r.doc_sim_model_2 for r in pairwise_results]
        ens_scores: List[float] = [r.doc_sim_ensemble for r in pairwise_results]

        avg_m1: float = sum(m1_scores) / len(m1_scores) if m1_scores else 0.0
        avg_m2: float = sum(m2_scores) / len(m2_scores) if m2_scores else 0.0
        avg_ens: float = sum(ens_scores) / len(ens_scores) if ens_scores else 0.0

        diffs: List[float] = [abs(s1 - s2) for s1, s2 in zip(m1_scores, m2_scores)]
        avg_diff: float = sum(diffs) / len(diffs) if diffs else 0.0

        # Writing statistics
        ws["A3"] = "Comparative Metric"
        ws["B3"] = "Value"
        self._style_header_row(ws, row_idx=3, max_col=2, bg_color="2874A6")

        stats: List[Tuple[str, str]] = [
            (f"Mean pair similarity ({self.model_1_name})", f"{avg_m1:.2%}"),
            (f"Mean pair similarity ({self.model_2_name})", f"{avg_m2:.2%}"),
            ("Mean pair similarity (Ensemble Model)", f"{avg_ens:.2%}"),
            ("Mean absolute divergence between models", f"{avg_diff:.2%}"),
            ("Number of detected suspicious sentences (above threshold)", str(len(matches))),
            (
                "Number of fragments with significant model divergence (>0.15)",
                str(sum(1 for m in matches if m.divergence > 0.15)),
            ),
        ]

        r_idx: int = 4
        for label, val in stats:
            ws.cell(row=r_idx, column=1, value=label).font = Font(bold=True)
            ws.cell(row=r_idx, column=2, value=val).alignment = Alignment(
                horizontal="center"
            )
            r_idx += 1

        # Conclusions section
        current: int = r_idx + 2
        ws.merge_cells(f"A{current}:F{current}")
        ws[f"A{current}"].value = "Conclusions from Model Fusion:"
        ws[f"A{current}"].font = Font(size=12, bold=True, color="1F497D")

        conclusions: List[str] = [
            "1. Complementarity of models: The model dedicated to the Polish language (ST-Polish-RoBERTa) is characterized by higher",
            "   sensitivity to complex inflectional structures and specific Polish phraseological expressions.",
            "2. The multilingual model (Multilingual-MiniLM) is more robust to rare borrowed vocabulary",
            "   as well as proper names and technical terms.",
            "3. Reduction of false positives: In the case of articles on the same topic, single models",
            "   can overestimate the similarity score due to identical keywords. The ensemble model smooths",
            "   out extreme deviations and provides more stable decision boundaries.",
            "4. Using a weighted average (e.g. alpha=0.55 in favor of the monolingual model) allows maintaining high",
            "   semantic precision while preserving the flexibility of the multilingual bi-encoder.",
        ]

        for line in conclusions:
            current += 1
            ws.cell(row=current, column=1, value=line).font = Font(
                italic=True, color="2C3E50"
            )

        self._auto_fit_columns(ws, max_len_cap=80)
